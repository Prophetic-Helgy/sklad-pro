"""Тесты миграции: импорт файла «Склад ПРО 7.1» и экспорт в его формат."""
import io
from datetime import datetime

from openpyxl import Workbook, load_workbook

NOM_HEADERS = ["r", "№", "Наименование", "Артикул", "Категория", "Ед. изм", "Кол-во (нач)",
               "Поставщик", "Примечание", "Приход (з)", "Расход (з)", "Списание (з)",
               "Приход (с)", "Расход (с)", "Списание (с)", "Приход", "Расход", "Списание"]
PARTNERS_HEADERS = ["r", "№", "Полное наименование", "Сокращенное наименование", "Юридический адрес",
                    "Почтовый адрес", "Тел/факс", "E-Mail", "ИНН", "ОГРН/ОГРНИП", "Примечание"]
IN_HEADERS = ["r", "Дата", "Наименование", "Артикул", "Категория", "Кол-во", "Ед. изм",
              "Поставщик", "Накладная", "Комментарий"]
OUT_HEADERS = ["r", "Дата", "Наименование", "Артикул", "Категория", "Кол-во", "Ед. изм",
               "Отпустил", "Получил", "Накладная", "Комментарий"]
SPIS_HEADERS = ["r", "Дата", "Наименование", "Артикул", "Категория", "Кол-во", "Ед. изм",
                "Поставщик", "Накладная (№, дата)", "Причина списания"]


def _put(ws, row: int, first_col: int, values) -> None:
    for i, v in enumerate(values):
        if v is not None:
            ws.cell(row=row, column=first_col + i, value=v)


def _build_old_file() -> bytes:
    """Синтетический файл 7.1 с данными."""
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Номенклатура")
    _put(ws, 2, 2, NOM_HEADERS)
    _put(ws, 3, 2, [None, None, "Болт М8", "B-M8", "Метизы", "шт", 10, "ООО «Строй»", None, None, None, None])
    _put(ws, 4, 2, [None, None, "Гайка М8", None, "Метизы", "шт", -3, None, None, None, None, None])
    _put(ws, 5, 2, [None, None, "Шайба 8", "Ш-8", None, "шт", 0, None, "Плоская", 5, None, None])
    _put(ws, 6, 2, [None, None, None, None, None, None, 7, None, None])  # пустое наименование → не считается

    ws = wb.create_sheet("Поставщики")
    _put(ws, 2, 2, PARTNERS_HEADERS)
    _put(ws, 3, 2, [None, None, "ООО «Строй»", "Строй", "г. Москва, ул. Зап.", None, "+7 495", None, "77001", "123", None])

    ws = wb.create_sheet("Реестр прихода")
    _put(ws, 3, 2, IN_HEADERS)
    _put(ws, 4, 2, [None, datetime(2026, 3, 1), "Болт М8", "B-M8", "Метизы", 10, "шт", "ООО «Строй»", "Н-100", None])
    _put(ws, 5, 2, [None, datetime(2026, 3, 1), "Гайка М8", None, "Метизы", 5, "шт", "ООО «Строй»", "Н-100", None])
    _put(ws, 6, 2, [None, datetime(2026, 3, 5), "Шайба 8", "Ш-8", None, 4, "шт", "ООО «Строй»", None, None])
    _put(ws, 7, 2, [None, "не дата", "Болт М8", None, None, 1, "шт", "ООО «Строй»", None, None])  # → пропуск
    _put(ws, 8, 2, [None, datetime(2026, 3, 5), "Болт М8", None, None, 0, "шт", "ООО «Строй»", None, None])  # → пропуск

    ws = wb.create_sheet("Реестр расхода")
    _put(ws, 3, 2, OUT_HEADERS)
    _put(ws, 4, 2, [None, datetime(2026, 3, 2), "Болт М8", "B-M8", "Метизы", 4, "шт", "Иван Петров", "ОО «Клиент»", None, None])
    _put(ws, 5, 2, [None, datetime(2026, 3, 1), "Гайка М8", None, "Метизы", 2, "шт", "Иван Петров", "ОО «Клиент»", None, None])

    ws = wb.create_sheet("Реестр списания")
    _put(ws, 3, 2, SPIS_HEADERS)
    _put(ws, 4, 2, [None, datetime(2026, 3, 10), "Болт М8", "B-M8", "Метизы", 1, "шт", "ООО «Строй»", None, "Брак"])

    ws = wb.create_sheet("Админ")
    ws.cell(row=4, column=3, value="Карточка предприятия")
    for r, label, val in ((6, "Полное наименование", "ООО «Старая компания»"),
                          (12, "ИНН/КПП", "7700000000/770001001"),
                          (14, "Руководитель", "Иванов И.И.")):
        ws.cell(row=r, column=3, value=label)
        ws.cell(row=r, column=4, value=val)
    _put(ws, 5, 7, [None, "Ф.И.О.", "Должность", "Логин", "Пароль"])
    _put(ws, 6, 7, [None, "Петров П.П.", None, "admin", "1234"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _section(report: dict, name: str) -> dict:
    return next(s for s in report["sections"] if s["name"] == name)


def test_import_old_file(client, auth):
    r = client.post(
        "/api/admin/migrate/import",
        files={"file": ("склад 7.1.xlsm", _build_old_file(),
                        "application/vnd.ms-excel.sheet.macroEnabled.12")},
        headers=auth,
    )
    assert r.status_code == 200, r.text
    rep = r.json()

    # Номенклатура: 3 товара созданы
    nom = _section(rep, "Номенклатура")
    assert nom["total"] == 3
    assert nom["created"] == 3

    # Поставщики: 1 из файла + 1 «Получил» из расхода
    par = _section(rep, "Поставщики")
    assert par["created"] == 1
    assert "ОО «Клиент»" in rep["created"]["partners"]

    # Сотрудники: admin найден по логину, пароль старый не меняет
    sotr = _section(rep, "Сотрудники")
    assert sotr["kept"] == 1
    assert any("пароль проигнорирован" in w for w in rep["warnings"])

    # Карточка предприятия: 3 поля
    comp = _section(rep, "Карточка предприятия")
    assert comp["fields"] == 3

    # Документы: начальные (приход 1, расход 1) + ручной (з) (1) + реестры (2+2+1) = 8
    assert rep["documents"]["total"] == 8
    assert rep["documents"]["by_type"] == {"Приход": 4, "Расход": 3, "Списание": 1}

    # Пропуски: не дата, нулевое кол-во
    rein = _section(rep, "Реестр прихода")
    reasons = [s["reason"] for s in rein["skipped_rows"]]
    assert any("дата" in s for s in reasons)
    assert any("количества" in s for s in reasons)
    assert rein["imported"] == 3

    # Сверка остатков: расхождений быть не должно
    assert rep["stock_check"]["mismatches"] == [], rep["stock_check"]

    # Остатки в базе: Болт 10+10−4−1=15, Гайка −3+5−2=0, Шайба 0+5+4=9
    r = client.get("/api/stock?size=500", headers=auth)
    assert r.status_code == 200
    body = r.json()
    rows = body["rows"] if "rows" in body else body
    stock_map: dict[str, float] = {}
    for row in rows:
        if isinstance(row, dict):
            name = row.get("item_name") or row.get("name")
            if name:
                stock_map[name] = (stock_map.get(name) or 0) + (row.get("qty") or row.get("qty_total") or 0)
    assert abs(stock_map.get("Болт М8", 0) - 15) < 0.001, stock_map
    assert abs(stock_map.get("Гайка М8", 0) - 0) < 0.001, stock_map
    assert abs(stock_map.get("Шайба 8", 0) - 9) < 0.001, stock_map


def test_import_rejected_and_non_admin(client, auth):
    r = client.post(
        "/api/admin/migrate/import",
        files={"file": ("bad.xlsm", b"not an excel file at all", "application/octet-stream")},
        headers=auth,
    )
    assert r.status_code == 400


def test_export_71_format(client, auth):
    r = client.get("/api/admin/migrate/export", headers=auth)
    assert r.status_code == 200
    assert r.headers["content-disposition"].startswith("attachment")
    wb = load_workbook(io.BytesIO(r.content), data_only=True)
    for sheet in ("Номенклатура", "Склад", "Поставщики", "Реестр прихода", "Реестр расхода", "Реестр списания", "Админ"):
        assert sheet in wb.sheetnames, f"нет листа {sheet}"
    # Номенклатура: заголовок на 2-й строке, данные с 3-й
    ws = wb["Номенклатура"]
    assert _norm_cell(ws.cell(row=2, column=4).value) == "Наименование"
    names = [ws.cell(row=r, column=4).value for r in range(3, ws.max_row + 1) if ws.cell(row=r, column=4).value]
    assert "Болт М8" in names
    # Реестр прихода: r=B, Дата=C, Наименование=D; строки с 4-й.
    # В реестре также лежат документы других тестов сессии — смотрим только свои товары:
    # начальные (Болт) + (з) (Шайба) + из реестров (Болт, Гайка, Шайба) = 5 строк
    ws = wb["Реестр прихода"]
    assert ws.cell(row=3, column=4).value == "Наименование"
    own = {"Болт М8", "Гайка М8", "Шайба 8"}
    own_rows = [r for r in range(4, ws.max_row + 1) if ws.cell(row=r, column=4).value in own]
    assert len(own_rows) == 5, [ws.cell(row=r, column=4).value for r in range(4, ws.max_row + 1)]
    assert "Болт М8" in [ws.cell(row=r, column=4).value for r in own_rows]
    # Склад: Остаток (факт) по Болту = 15
    ws = wb["Склад"]
    for r in range(3, ws.max_row + 1):
        if ws.cell(row=r, column=4).value == "Болт М8":
            assert abs(float(ws.cell(row=r, column=9).value) - 15) < 0.001
    # Админ: карточка предприятия
    ws = wb["Админ"]
    assert ws.cell(row=6, column=4).value == "ООО «Старая компания»"


def _norm_cell(v) -> str:
    return str(v).strip() if v is not None else ""
