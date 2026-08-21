"""Импорт/экспорт файла старого «Склад ПРО 7.1» (.xlsm/.xlsx).

Импорт: листы «Номенклатура», «Поставщики», «Реестр прихода/расхода/списания»
и «Админ» (карточка предприятия + сотрудники Sotrud) переносятся в новую
модель. Всё — в одной транзакции: либо целиком, либо откат. Функция
возвращает отчёт: что перенесено, что пропущено и почему, сверка остатков.

Экспорт: файл с той же структурой листов, что в 7.1 (значения, без формул и макросов).
"""
import io
import re
import secrets
from datetime import date, datetime, timedelta

from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from . import core, posting
from .models import (
    DOC_IN, DOC_OUT, DOC_WRITEOFF,
    ST_DRAFT, ST_POSTED,
    Category, Company, Document, DocumentItem, Item, Partner, Setting, Stock, Unit, User, Warehouse,
)
from .security import hash_password

SHEET_NOM = "Номенклатура"
SHEET_PARTNERS = "Поставщики"
SHEET_REESTR_IN = "Реестр прихода"
SHEET_REESTR_OUT = "Реестр расхода"
SHEET_REESTR_SPIS = "Реестр списания"
SHEET_ADMIN = "Админ"

# Подпись карточки предприятия (колонка C листа «Админ») → поле модели Company
COMPANY_LABELS = {
    "полное наименование": "full_name",
    "сокращенное наименование": "short_name",
    "юр. адрес": "legal_address",
    "юридический адрес": "legal_address",
    "почт. адрес": "postal_address",
    "почтовый адрес": "postal_address",
    "телефон/факс": "phone",
    "телефон / факс": "phone",
    "e-mail": "email",
    "email": "email",
    "инн/кпп": "inn_kpp",
    "инн / кпп": "inn_kpp",
    "огрн": "ogrn",
    "руководитель": "director",
    "банк": "bank",
    "бик": "bik",
    "расчетный счет": "account",
    "расчётный счет": "account",
    "корреспондентский счет": "corr_account",
    "корр. счет": "corr_account",
}

DOC_TYPE_LABEL = {DOC_IN: "Приход", DOC_OUT: "Расход", DOC_WRITEOFF: "Списание"}


# ================= Чтение ячеек =================

def _norm(v) -> str:
    """Ячейка → строка без пробков по краям (None → '')."""
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v) and abs(v) < 1e15:
        return str(int(v))
    return str(v).strip()


def _num(v) -> float | None:
    """Ячейка → число (None/пустое → 0.0; нераспознаваемо → None)."""
    if v is None or _norm(v) == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = _norm(v).replace(" ", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _parse_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = _norm(v)
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _find_header_row(ws, marker: str, max_row: int = 12) -> int | None:
    """Строка заголовка таблицы — ищем ячейку с точным текстом маркера."""
    for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=25):
        for c in row:
            if c.value is not None and _norm(c.value) == marker:
                return c.row
    return None


def _header_map(ws, header_row: int) -> dict[str, int]:
    """Заголовки строки → номера колонок (имена без пробелов)."""
    m: dict[str, int] = {}
    for c in range(1, 30):
        name = _norm(ws.cell(row=header_row, column=c).value)
        if name:
            m[name] = c
    return m


def _cell(ws, row: int, col_map: dict[str, int], key: str):
    col = col_map.get(key)
    return None if col is None else ws.cell(row=row, column=col).value


def _slug(name: str) -> str:
    """Латиница/кириллица → логин/арт: нижний регистр, не-буквы → '-', без хвостов."""
    s = re.sub(r"[^a-zа-яё0-9]+", "-", name.lower()).strip("-")
    return s[:40] or "x"


# ================= Импорт =================

def import_xlsm(db: Session, data: bytes, user_id: int | None = None) -> dict:
    """Перенос данных из файла 7.1. Возвращает отчёт (dict для JSON)."""
    try:
        return _import_inner(db, data, user_id)
    except Exception:
        db.rollback()
        raise


def _import_inner(db: Session, data: bytes, user_id: int | None) -> dict:
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as e:
        raise ValueError(f"Файл не читается как Excel: {e}") from e

    report: dict = {
        "ok": True,
        "sections": [],
        "created": {"categories": [], "units": [], "partners": [], "items": [], "users": []},
        "documents": {"total": 0, "by_type": {}},
        "stock_check": {"checked": 0, "mismatches": []},
        "warnings": [],
    }

    # --- Склад по умолчанию (в 7.1 одна глобальная «книга») ---
    wh_id = _default_warehouse_id(db)

    # --- Контроль отрицательных остатков на время переноса ---
    # История 7.1 не гарантирует «приход раньше расхода» — включаем, вернём после.
    neg = db.get(Setting, "allow_negative_stock")
    neg_old = neg.value if neg else None
    if neg is None:
        neg = Setting(key="allow_negative_stock")
        db.add(neg)
    neg.value = "true"
    db.flush()

    try:
        # Кэши (имя → объект) на время импорта
        units: dict[str, Unit] = {}
        cats: dict[str, Category] = {}
        partners: dict[str, Partner] = {}
        items: dict[str, Item] = {}
        users: dict[str, User] = {}
        default_unit: Unit | None = None

        def ensure_unit() -> Unit:
            """Ед. изм. по умолчанию (в модели unit_id обязателен, в 7.1 может быть пустой)."""
            nonlocal default_unit
            if default_unit is None:
                default_unit = (
                    unit_of("шт")
                    or db.scalar(select(Unit).order_by(Unit.id).limit(1))
                )
                if default_unit is None:
                    default_unit = Unit(name="шт")
                    db.add(default_unit)
                    db.flush()
            return default_unit
        sku_taken: set[str] = {
            s.lower() for (s,) in db.execute(select(Item.sku))
        }

        def unit_of(name: str) -> Unit | None:
            name = _norm(name)
            if not name:
                return None
            key = name.lower()
            if key in units:
                return units[key]
            u = db.scalar(select(Unit).where(func.ci_lower(Unit.name) == key))
            if u is None:
                u = Unit(name=name)
                db.add(u)
                db.flush()
                report["created"]["units"].append(name)
            units[key] = u
            return u

        def cat_of(name: str) -> Category | None:
            name = _norm(name)
            if not name:
                return None
            key = name.lower()
            if key in cats:
                return cats[key]
            c = db.scalar(select(Category).where(func.ci_lower(Category.name) == key))
            if c is None:
                c = Category(name=name)
                db.add(c)
                db.flush()
                report["created"]["categories"].append(name)
            cats[key] = c
            return c

        def partner_of(name: str, extra: dict | None = None) -> Partner | None:
            name = _norm(name)
            if not name:
                return None
            key = name.lower()
            if key in partners:
                return partners[key]
            p = db.scalar(select(Partner).where(func.ci_lower(Partner.full_name) == key))
            if p is None:
                p = Partner(full_name=name, **dict(extra or {}))
                db.add(p)
                db.flush()
                report["created"]["partners"].append(name)
            partners[key] = p
            return p

        def unique_sku(name: str, sku: str) -> str:
            """Уникальный артикул (в 7.1 уникальность не контролировалась)."""
            base = _norm(sku) or _slug(name)
            cand, n = base, 2
            while cand.lower() in sku_taken:
                cand = f"{base}-{n}"
                n += 1
            sku_taken.add(cand.lower())
            return cand

        def item_of(name: str, sku: str = "", unit: str = "",
                    cat: str = "", supplier: str = "", note: str = "") -> Item | None:
            name = _norm(name)
            if not name:
                return None
            key = name.lower()
            if key in items:
                it = items[key]
            else:
                it = db.scalar(select(Item).where(func.ci_lower(Item.name) == key))
                if it is None:
                    u = unit_of(unit) or ensure_unit()
                    c = cat_of(cat)
                    p = partner_of(supplier) if _norm(supplier) else None
                    it = Item(
                        name=name,
                        sku=unique_sku(name, sku),
                        unit_id=u.id,
                        category_id=c.id if c else None,
                        default_partner_id=p.id if p else None,
                        notes=_norm(note) or None,
                    )
                    db.add(it)
                    db.flush()
                    report["created"]["items"].append(name)
                else:
                    # Существующий товар: добираем только пустые поля
                    if _norm(cat) and it.category_id is None:
                        c = cat_of(cat)
                        if c:
                            it.category_id = c.id
                    if _norm(supplier) and it.default_partner_id is None:
                        p = partner_of(supplier)
                        if p:
                            it.default_partner_id = p.id
                    if _norm(note) and not it.notes:
                        it.notes = _norm(note)
                items[key] = it
            return it

        def user_of(full_name: str, username: str = "", make: bool = False) -> User | None:
            full_name = _norm(full_name)
            username = _norm(username)
            key = (full_name or username).lower()
            if key in users:
                return users[key]
            u = None
            if full_name:
                u = db.scalar(select(User).where(func.ci_lower(User.full_name) == full_name.lower()))
            if u is None and username:
                u = db.scalar(select(User).where(func.ci_lower(User.username) == username.lower()))
            if u is None and make:
                uname = _slug(full_name or username)
                n, base = 2, uname
                while db.scalar(select(User).where(func.ci_lower(User.username) == base.lower())):
                    base = f"{uname}{n}"
                    n += 1
                password = secrets.token_urlsafe(8)
                u = User(
                    username=base,
                    password_hash=hash_password(password),
                    full_name=full_name or None,
                    role="admin" if (username or base).lower() == "admin" else "user",
                    active=True,
                )
                db.add(u)
                db.flush()
                report["created"]["users"].append({"username": base, "password": password})
            users[key] = u
            return u

        # ---------- 1. Поставщики ----------
        sec = _section("Поставщики")
        ws = wb[SHEET_PARTNERS] if SHEET_PARTNERS in wb.sheetnames else None
        if ws is None:
            sec["skipped_rows"].append({"row": 0, "reason": "лист отсутствует в файле"})
        else:
            hr = _find_header_row(ws, "Полное наименование")
            if hr is None:
                sec["skipped_rows"].append({"row": 0, "reason": "заголовок таблицы не найден"})
            else:
                hm = _header_map(ws, hr)
                for r in range(hr + 1, ws.max_row + 1):
                    name = _norm(_cell(ws, r, hm, "Полное наименование"))
                    if not name:
                        continue
                    sec["total"] += 1
                    extra = {
                        "short_name": _norm(_cell(ws, r, hm, "Сокращенное наименование")) or None,
                        "legal_address": _norm(_cell(ws, r, hm, "Юридический адрес")) or None,
                        "postal_address": _norm(_cell(ws, r, hm, "Почтовый адрес")) or None,
                        "phone": _norm(_cell(ws, r, hm, "Тел/факс")) or None,
                        "email": _norm(_cell(ws, r, hm, "E-Mail")) or None,
                        "inn": _norm(_cell(ws, r, hm, "ИНН")) or None,
                        "ogrn": _norm(_cell(ws, r, hm, "ОГРН/ОГРНИП")) or None,
                        "notes": _norm(_cell(ws, r, hm, "Примечание")) or None,
                    }
                    before = len(report["created"]["partners"])
                    partner_of(name, extra)
                    if len(report["created"]["partners"]) > before:
                        sec["created"] += 1
                    else:
                        sec["kept"] += 1
        report["sections"].append(sec)

        # ---------- 2. Номенклатура ----------
        sec = _section("Номенклатура")
        ws = wb[SHEET_NOM] if SHEET_NOM in wb.sheetnames else None
        open_in: list[tuple[Item, float]] = []    # начальные остатки > 0
        open_out: list[tuple[Item, float]] = []    # начальные остатки < 0
        manual_in: list[tuple[Item, float]] = []   # Приход (з)
        manual_out: list[tuple[Item, float]] = []  # Расход (з)
        manual_spis: list[tuple[Item, float]] = []  # Списание (з)
        expected: dict[str, float] = {}            # имя → ожидаемый остаток
        if ws is None:
            sec["skipped_rows"].append({"row": 0, "reason": "лист отсутствует в файле"})
        else:
            hr = _find_header_row(ws, "Наименование")
            if hr is None:
                sec["skipped_rows"].append({"row": 0, "reason": "заголовок таблицы не найден"})
            else:
                hm = _header_map(ws, hr)
                for r in range(hr + 1, ws.max_row + 1):
                    name = _norm(_cell(ws, r, hm, "Наименование"))
                    if not name:
                        continue
                    sec["total"] += 1
                    before = len(report["created"]["items"])
                    it = item_of(
                        name,
                        sku=_norm(_cell(ws, r, hm, "Артикул")),
                        unit=_norm(_cell(ws, r, hm, "Ед. изм")),
                        cat=_norm(_cell(ws, r, hm, "Категория")),
                        supplier=_norm(_cell(ws, r, hm, "Поставщик")),
                        note=_norm(_cell(ws, r, hm, "Примечание")),
                    )
                    if it is None:
                        sec["skipped_rows"].append({"row": r, "reason": "не удалось создать товар"})
                        continue
                    if len(report["created"]["items"]) > before:
                        sec["created"] += 1
                    else:
                        sec["kept"] += 1
                    qty_open = _num(_cell(ws, r, hm, "Кол-во (нач)")) or 0.0
                    z_in = _num(_cell(ws, r, hm, "Приход (з)")) or 0.0
                    z_out = _num(_cell(ws, r, hm, "Расход (з)")) or 0.0
                    z_spis = _num(_cell(ws, r, hm, "Списание (з)")) or 0.0
                    if qty_open > 0:
                        open_in.append((it, qty_open))
                    elif qty_open < 0:
                        open_out.append((it, -qty_open))
                    if z_in > 0:
                        manual_in.append((it, z_in))
                    if z_out > 0:
                        manual_out.append((it, z_out))
                    if z_spis > 0:
                        manual_spis.append((it, z_spis))
                    expected[name.lower()] = qty_open + z_in - z_out - z_spis
        report["sections"].append(sec)

        # ---------- 3. Реестры (история фактов → документы) ----------
        reg_docs: list[dict] = []  # спецификации документов: {type, date, ...}

        def parse_reestr(sheet: str, type_: str, partner_key: str, extra_keys: list[str] | None = None) -> None:
            sec = _section(sheet)
            ws = wb[sheet] if sheet in wb.sheetnames else None
            if ws is None:
                sec["skipped_rows"].append({"row": 0, "reason": "лист отсутствует в файле"})
                report["sections"].append(sec)
                return
            hr = _find_header_row(ws, "Дата")
            if hr is None:
                sec["skipped_rows"].append({"row": 0, "reason": "заголовок таблицы не найден"})
                report["sections"].append(sec)
                return
            hm = _header_map(ws, hr)
            groups: dict[tuple, dict] = {}
            for r in range(hr + 1, ws.max_row + 1):
                name = _norm(_cell(ws, r, hm, "Наименование"))
                d = _parse_date(_cell(ws, r, hm, "Дата"))
                qty = _num(_cell(ws, r, hm, "Кол-во"))
                if not name and d is None and (qty or 0) == 0:
                    continue
                sec["total"] += 1
                if d is None:
                    sec["skipped_rows"].append({"row": r, "reason": f"не распознана дата: «{_norm(_cell(ws, r, hm, 'Дата'))}»"})
                    continue
                if not name:
                    sec["skipped_rows"].append({"row": r, "reason": "пустое наименование"})
                    continue
                if qty is None or qty == 0:
                    sec["skipped_rows"].append({"row": r, "reason": "нет количества (пусто или 0)"})
                    continue
                partner = _norm(_cell(ws, r, hm, partner_key))
                extra = tuple(_norm(_cell(ws, r, hm, k)) for k in (extra_keys or []))
                key = (d, partner, *extra, _norm(_cell(ws, r, hm, "Накладная")) if "Накладная" in hm else "")
                g = groups.get(key)
                if g is None:
                    g = {
                        "type": type_, "date": d, "partner": partner,
                        **{k: _norm(_cell(ws, r, hm, k)) for k in (extra_keys or [])},
                        "doc_ref": _norm(_cell(ws, r, hm, "Накладная")) if "Накладная" in hm else (
                            _norm(_cell(ws, r, hm, "Накладная (№, дата)")) if "Накладная (№, дата)" in hm else ""),
                        "reason": _norm(_cell(ws, r, hm, "Причина списания")) if "Причина списания" in hm else "",
                        "lines": [],
                    }
                    groups[key] = g
                g["lines"].append((
                    name,
                    _norm(_cell(ws, r, hm, "Артикул")),
                    _norm(_cell(ws, r, hm, "Ед. изм")),
                    qty,
                    _norm(_cell(ws, r, hm, "Комментарий")) or _norm(_cell(ws, r, hm, "Примечание")),
                ))
                sec["imported"] += 1
                # Ожидаемый остаток: приход +, расход/списание −
                sign = 1 if type_ == DOC_IN else -1
                expected[name.lower()] = expected.get(name.lower(), 0.0) + sign * qty
            for g in groups.values():
                reg_docs.append(g)
            sec["documents"] = len(groups)
            report["sections"].append(sec)

        parse_reestr(SHEET_REESTR_IN, DOC_IN, "Поставщик")
        parse_reestr(SHEET_REESTR_OUT, DOC_OUT, "Получил", extra_keys=["Отпустил"])
        parse_reestr(SHEET_REESTR_SPIS, DOC_WRITEOFF, "Поставщик")

        # ---------- 4. Сотрудники (Sotrud, лист «Админ») ----------
        sec = _section("Сотрудники")
        ws = wb[SHEET_ADMIN] if SHEET_ADMIN in wb.sheetnames else None
        if ws is None:
            sec["skipped_rows"].append({"row": 0, "reason": "лист отсутствует в файле"})
        else:
            hr = _find_header_row(ws, "Логин", max_row=20)
            if hr is None:
                pass  # без Sotrud — не ошибка (в файле может быть только карточка)
            else:
                hm = _header_map(ws, hr)
                for r in range(hr + 1, ws.max_row + 1):
                    login = _norm(_cell(ws, r, hm, "Логин"))
                    full_name = _norm(_cell(ws, r, hm, "Ф.И.О."))
                    if not login and not full_name:
                        continue
                    sec["total"] += 1
                    password = _norm(_cell(ws, r, hm, "Пароль"))
                    u = user_of(full_name, login, make=bool(login))
                    if u is None:
                        sec["skipped_rows"].append({"row": r, "reason": f"не удалось создать пользователя «{login or full_name}»"})
                    elif (u.username or "").lower() == "admin" and u.role != "admin":
                        u.role = "admin"
                        sec["updated"] += 1
                    else:
                        sec["kept"] += 1
                    if password:
                        report["warnings"].append(
                            f"Пользователь «{u.username}»: старый пароль проигнорирован (оставлен текущий)"
                        )
        report["sections"].append(sec)

        # ---------- 5. Карточка предприятия ----------
        sec = _section("Карточка предприятия")
        if ws is not None:
            comp = db.get(Company, 1)
            if comp is None:
                comp = Company(id=1)
                db.add(comp)
            for r in range(1, ws.max_row + 1):
                label = _norm(ws.cell(row=r, column=3).value).lower()
                field = COMPANY_LABELS.get(label)
                if field is None:
                    continue
                value = _norm(ws.cell(row=r, column=4).value)
                if value and value.lower() not in ("=today()",):
                    setattr(comp, field, value)
                    sec["fields"] += 1
        report["sections"].append(sec)

        # ---------- 6. Проведение документов ----------
        # Остатки товаров из файла до проведения (для сверки при доимпорте в непустую базу)
        pre_stock: dict[str, float] = {}
        for s in db.scalars(select(Stock)):
            it = db.get(Item, s.item_id)
            if it and it.name.lower() in expected:
                pre_stock[it.name.lower()] = pre_stock.get(it.name.lower(), 0.0) + float(s.qty or 0)
        # 6.1 начальные остатки (дата — за день до старшей записи реестров)
        all_dates = [g["date"] for g in reg_docs]
        base_date = (min(all_dates) - timedelta(days=1)) if all_dates else date.today()
        docs_made = 0

        def make_doc(type_: str, d: date, lines: list[tuple[Item, float]], **kw) -> None:
            nonlocal docs_made
            lines = [(it, q) for it, q in lines if q]
            if not lines:
                return
            doc = Document(type=type_, date=d, status=ST_DRAFT, **kw)
            db.add(doc)
            db.flush()
            for it, q in lines:
                doc.items.append(DocumentItem(
                    item_id=it.id,
                    qty=q,
                    unit_snapshot=it.unit_name if it.unit_id else None,
                    warehouse_to_id=wh_id if type_ == DOC_IN else None,
                    warehouse_from_id=None if type_ == DOC_IN else wh_id,
                ))
            posting.apply_doc(db, doc)
            doc.number = core.next_document_number(db, doc.type, d.year)
            doc.status = ST_POSTED
            doc.posted_at = datetime.utcnow()
            db.flush()
            docs_made += 1
            report["documents"]["by_type"][DOC_TYPE_LABEL[type_]] = (
                report["documents"]["by_type"].get(DOC_TYPE_LABEL[type_], 0) + 1
            )

        make_doc(DOC_IN, base_date, open_in, notes="Начальные остатки (перенос из 7.1)")
        make_doc(DOC_OUT, base_date, open_out, notes="Отрицательные начальные остатки (перенос из 7.1)")
        make_doc(DOC_IN, base_date, manual_in, notes="Ручные корректировки «(з)» (перенос из 7.1)")
        make_doc(DOC_OUT, base_date, manual_out, notes="Ручные корректировки «(з)» (перенос из 7.1)")
        make_doc(DOC_WRITEOFF, base_date, manual_spis, reason="Ручная корректировка (перенос из 7.1)")

        # 6.2 документы из реестров (в порядке дат)
        for g in sorted(reg_docs, key=lambda x: x["date"]):
            kw: dict = {}
            if g["type"] == DOC_IN:
                p = partner_of(g["partner"]) if g["partner"] else None
                if p:
                    kw["partner_id"] = p.id
            elif g["type"] == DOC_OUT:
                # partner_key для расхода — «Получил», хранится в g["partner"]
                p = partner_of(g["partner"]) if g["partner"] else None
                if p:
                    kw["partner_id"] = p.id
                u = user_of(g["Отпустил"], make=True) if g.get("Отпустил") else None
                if u:
                    kw["out_by_id"] = u.id
            else:  # Списание
                p = partner_of(g["partner"]) if g["partner"] else None
                if p:
                    kw["partner_id"] = p.id
                if g["reason"]:
                    kw["reason"] = g["reason"]
            if g["doc_ref"]:
                kw["doc_ref"] = g["doc_ref"][:50]
            lines = []
            for (name, sku, unit, qty, note) in g["lines"]:
                it = item_of(name, sku=sku, unit=unit)
                if it is None:
                    continue
                lines.append((it, qty))
            make_doc(g["type"], g["date"], lines, **kw)
        report["documents"]["total"] = docs_made

        # ---------- 7. Сверка остатков ----------
        actual: dict[str, float] = {}
        for s in db.scalars(select(Stock)):
            it = db.get(Item, s.item_id)
            if it:
                actual[it.name.lower()] = actual.get(it.name.lower(), 0.0) + float(s.qty or 0)
        # Сверяем только товары из файла: остаток после = был до + количество из файла
        for name in sorted(expected):
            exp, act = expected[name] + pre_stock.get(name, 0.0), actual.get(name, 0.0)
            if abs(exp - act) > 0.001:
                report["stock_check"]["mismatches"].append({
                    "item": next((n for n in items if n.lower() == name), name.title()),
                    "expected": round(exp, 3),
                    "imported": round(act, 3),
                    "diff": round(act - exp, 3),
                })
        report["stock_check"]["checked"] = len(expected)

        # ---------- 8. Финал ----------
        neg.value = neg_old or "false"
        core.audit(db, user_id, "import", "migration", None, {
            "items": len(report["created"]["items"]),
            "documents": docs_made,
            "sections": len(report["sections"]),
        })
        db.commit()
    finally:
        wb.close()
    return report


def _section(name: str) -> dict:
    return {"name": name, "total": 0, "created": 0, "kept": 0,
            "updated": 0, "imported": 0, "documents": 0, "fields": 0, "skipped_rows": []}


def _default_warehouse_id(db: Session) -> int:
    wh_id = core.default_warehouse_id(db)
    if wh_id:
        return wh_id
    wh = db.scalar(select(Warehouse).order_by(Warehouse.id).limit(1))
    if wh is None:
        wh = Warehouse(name="Основной")
        db.add(wh)
        db.flush()
    return wh.id


# ================= Экспорт в формат 7.1 =================

def export_71(db: Session) -> tuple[bytes, str]:
    """Файл с той же структурой листов, что в 7.1 (значения, без формул)."""
    wb = Workbook()
    wb.remove(wb.active)

    items = list(db.scalars(select(Item).where(Item.active.is_(True)).order_by(Item.id)))
    partners = {p.id: p for p in db.scalars(select(Partner))}
    units = {u.id: u for u in db.scalars(select(Unit))}
    cats = {c.id: c for c in db.scalars(select(Category))}
    users = {u.id: u for u in db.scalars(select(User))}
    stock: dict[int, float] = {}
    for s in db.scalars(select(Stock)):
        stock[s.item_id] = stock.get(s.item_id, 0.0) + float(s.qty or 0)

    # Подтоки по товарам по проведённым документам
    flow_in: dict[int, float] = {}
    flow_out: dict[int, float] = {}
    flow_spis: dict[int, float] = {}
    reestr_in, reestr_out, reestr_spis = [], [], []
    for doc in db.scalars(select(Document).where(Document.status == ST_POSTED).order_by(Document.date)):
        if doc.type not in (DOC_IN, DOC_OUT, DOC_WRITEOFF):
            continue
        rows = []
        for li in doc.items:
            it = db.get(Item, li.item_id)
            if it is None:
                continue
            partner = partners.get(doc.partner_id)
            d = datetime(doc.date.year, doc.date.month, doc.date.day)
            if doc.type == DOC_IN:
                flow_in[it.id] = flow_in.get(it.id, 0.0) + float(li.qty)
                rows.append((d, it, float(li.qty), _unit_name(it, units), partner.full_name if partner else "", doc.doc_ref or "", li.note or ""))
            elif doc.type == DOC_OUT:
                flow_out[it.id] = flow_out.get(it.id, 0.0) + float(li.qty)
                out_by = users.get(doc.out_by_id)
                rec = partner.full_name if partner else (doc.received_by or "")
                rows.append((d, it, float(li.qty), _unit_name(it, units), out_by.full_name if out_by else "", rec, doc.doc_ref or "", li.note or ""))
            else:
                flow_spis[it.id] = flow_spis.get(it.id, 0.0) + float(li.qty)
                rows.append((d, it, float(li.qty), _unit_name(it, units), partner.full_name if partner else "", doc.doc_ref or "", doc.reason or li.note or ""))
        if doc.type == DOC_IN:
            reestr_in.extend(rows)
        elif doc.type == DOC_OUT:
            reestr_out.extend(rows)
        else:
            reestr_spis.extend(rows)

    # ---------- Номенклатура ----------
    ws = wb.create_sheet(SHEET_NOM)
    _put_header(ws, 2, 2, ["r", "№", "Наименование", "Артикул", "Категория", "Ед. изм",
                           "Кол-во (нач)", "Поставщик", "Примечание",
                           "Приход (з)", "Расход (з)", "Списание (з)",
                           "Приход (с)", "Расход (с)", "Списание (с)",
                           "Приход", "Расход", "Списание"])
    r = 3
    for n, it in enumerate(items, 1):
        fin, fout, fsp = flow_in.get(it.id, 0.0), flow_out.get(it.id, 0.0), flow_spis.get(it.id, 0.0)
        vals = ["", n, it.name, it.sku, cats[it.category_id].name if it.category_id and it.category_id in cats else "",
                units[it.unit_id].name if it.unit_id and it.unit_id in units else "",
                0, partners[it.default_partner_id].full_name if it.default_partner_id and it.default_partner_id in partners else "",
                it.notes or "",
                0, 0, 0, fin, fout, fsp, fin, fout, fsp]
        for c, v in enumerate(vals, 2):
            ws.cell(row=r, column=c, value=v if v != "" else None)
        r += 1

    # ---------- Склад ----------
    ws = wb.create_sheet("Склад")
    _put_header(ws, 2, 3, ["№", "Наименование", "Артикул", "Категория", "Ед. изм",
                           "Кол-во (нач)", "Остаток (факт)", "Поставщик"])
    r = 3
    for n, it in enumerate(items, 1):
        vals = [n, it.name, it.sku,
                cats[it.category_id].name if it.category_id and it.category_id in cats else "",
                units[it.unit_id].name if it.unit_id and it.unit_id in units else "",
                0, round(stock.get(it.id, 0.0), 3),
                partners[it.default_partner_id].full_name if it.default_partner_id and it.default_partner_id in partners else ""]
        for c, v in enumerate(vals, 3):
            ws.cell(row=r, column=c, value=v if v != "" else None)
        r += 1

    # ---------- Поставщики ----------
    ws = wb.create_sheet(SHEET_PARTNERS)
    _put_header(ws, 2, 2, ["r", "№", "Полное наименование", "Сокращенное наименование", "Юридический адрес",
                           "Почтовый адрес", "Тел/факс", "E-Mail", "ИНН", "ОГРН/ОГРНИП", "Примечание"])
    r = 3
    for n, p in enumerate(sorted(partners.values(), key=lambda x: x.id), 1):
        vals = ["", n, p.full_name, p.short_name or None, p.legal_address or None, p.postal_address or None,
                p.phone or None, p.email or None, p.inn or None, p.ogrn or None, p.notes or None]
        for c, v in enumerate(vals, 2):
            ws.cell(row=r, column=c, value=v)
        r += 1

    # ---------- Реестры ----------
    def _reestr_sheet(name: str, headers: list[str], rows: list) -> None:
        ws = wb.create_sheet(name)
        _put_header(ws, 3, 2, ["r"] + headers)
        r = 4
        for row in rows:
            for c, v in enumerate(row, 2):
                cell = ws.cell(row=r, column=c, value=v if v not in ("", None) else None)
                if isinstance(v, datetime):
                    cell.number_format = "DD.MM.YYYY"
            r += 1

    def _cat(it: Item) -> str:
        return cats[it.category_id].name if it.category_id and it.category_id in cats else ""

    _reestr_sheet(
        SHEET_REESTR_IN,
        ["Дата", "Наименование", "Артикул", "Категория", "Кол-во", "Ед. изм", "Поставщик", "Накладная", "Комментарий"],
        [("", d, it.name, it.sku, _cat(it), q, un, pn, ref, note)
         for (d, it, q, un, pn, ref, note) in reestr_in],
    )
    _reestr_sheet(
        SHEET_REESTR_OUT,
        ["Дата", "Наименование", "Артикул", "Категория", "Кол-во", "Ед. изм", "Отпустил", "Получил", "Накладная", "Комментарий"],
        [("", d, it.name, it.sku, _cat(it), q, un, ob, rec, ref, note)
         for (d, it, q, un, ob, rec, ref, note) in reestr_out],
    )
    _reestr_sheet(
        SHEET_REESTR_SPIS,
        ["Дата", "Наименование", "Артикул", "Категория", "Кол-во", "Ед. изм", "Поставщик", "Накладная (№, дата)", "Причина списания"],
        [("", d, it.name, it.sku, _cat(it), q, un, pn, ref, note)
         for (d, it, q, un, pn, ref, note) in reestr_spis],
    )

    # ---------- Админ (карточка + Sotrud) ----------
    ws = wb.create_sheet(SHEET_ADMIN)
    ws.cell(row=4, column=3, value="Карточка предприятия")
    comp = db.get(Company, 1)
    if comp is not None:
        for r_, label, val in (
            (6, "Полное наименование", comp.full_name),
            (7, "Сокращенное наименование", comp.short_name),
            (8, "Юридический адрес", comp.legal_address),
            (9, "Почтовый адрес", comp.postal_address),
            (10, "Телефон/факс", comp.phone),
            (11, "E-mail", comp.email),
            (12, "ИНН/КПП", comp.inn_kpp),
            (13, "ОГРН", comp.ogrn),
            (14, "Руководитель", comp.director),
            (15, "Банковские реквизиты", None),
            (16, "Банк", comp.bank),
            (17, "БИК", comp.bik),
            (18, "Расчетный счет", comp.account),
            (19, "Корреспондентский счет", comp.corr_account),
        ):
            ws.cell(row=r_, column=3, value=label)
            if val:
                ws.cell(row=r_, column=4, value=val)
    for c, h in enumerate(["r", "Ф.И.О.", "Должность", "Логин", "Пароль"], 7):
        ws.cell(row=5, column=c, value=h)
    for i, u in enumerate(sorted(users.values(), key=lambda x: x.id)):
        ws.cell(row=6 + i, column=8, value=u.full_name)
        ws.cell(row=6 + i, column=10, value=u.username)

    buf = io.BytesIO()
    wb.save(buf)
    stamp = date.today().strftime("%Y-%m-%d")
    return buf.getvalue(), f"Склад ПРО 7.1_{stamp}.xlsx"


def _put_header(ws, row: int, first_col: int, headers: list[str]) -> None:
    for i, h in enumerate(headers):
        ws.cell(row=row, column=first_col + i, value=h)


def _unit_name(it: Item, units: dict[int, Unit]) -> str:
    return units[it.unit_id].name if it.unit_id and it.unit_id in units else ""
