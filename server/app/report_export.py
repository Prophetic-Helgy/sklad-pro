"""Отчёт (общая структура) и экспорт в Excel (.xlsx) и PDF.

Отчёт — это заголовок, период, метаданные колонок и строки-значения;
фронтенд и экспорт строятся от одной и той же структуры.
"""
import io
from dataclasses import dataclass, field

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

try:  # PDF-часть опциональна: без reportlab Excel продолжает работать
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Spacer, Table, TableStyle

    _PDF_FONT = "Helvetica"
    for _cand in (r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\tahoma.ttf"):
        try:
            pdfmetrics.registerFont(TTFont("SkladFont", _cand))
            _PDF_FONT = "SkladFont"
            break
        except Exception:
            continue
    _HAS_PDF = True
except ImportError:  # pragma: no cover
    _HAS_PDF = False


@dataclass
class Report:
    """Собранный отчёт: колонки + строки + итоговая строка."""

    title: str
    period: str | None = None
    columns: list = field(default_factory=list)   # [{key, title, align, kind, width}]
    rows: list = field(default_factory=list)      # [[значение по каждой колонке]]
    totals: list | None = None                    # строка «Итого» (или None)


def col(key: str, title: str, kind: str = "text", align: str | None = None, width: int | None = None) -> dict:
    """Описание колонки отчёта. kind: text | number | money | percent."""
    return {"key": key, "title": title, "kind": kind, "align": align, "width": width}


def _is_num(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


# ---------------------------------------------------------------- Excel
def build_excel(rep: Report) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = (rep.title or "Отчёт")[:31]

    thin = Side(style="thin", color="9CA3AF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.cell(row=1, column=1, value=rep.title).font = Font(bold=True, size=14)
    if rep.period:
        ws.cell(row=2, column=1, value=rep.period).font = Font(color="6B7280")

    header_row = 4
    for j, c in enumerate(rep.columns, start=1):
        cell = ws.cell(row=header_row, column=j, value=c["title"])
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    r = header_row + 1
    for row in rep.rows:
        for j, (c, v) in enumerate(zip(rep.columns, row), start=1):
            cell = ws.cell(row=r, column=j, value=v)
            cell.border = border
            if _is_num(v):
                cell.number_format = _excel_format(c["kind"])
                cell.alignment = Alignment(horizontal="right")
        r += 1

    if rep.totals is not None:
        top = Border(top=Side(style="double", color="374151"))
        for j, (c, v) in enumerate(zip(rep.columns, rep.totals), start=1):
            cell = ws.cell(row=r, column=j, value=v)
            cell.font = Font(bold=True)
            cell.border = top
            if _is_num(v):
                cell.number_format = _excel_format(c["kind"])
                cell.alignment = Alignment(horizontal="right")

    # Ширина колонок: под самую длинную ячейку (кириллица шире латиницы)
    for j, c in enumerate(rep.columns, start=1):
        width = len(c["title"]) * 1.15
        for row in rep.rows:
            v = row[j - 1]
            if v is None:
                continue
            width = max(width, len(str(v)) * (1.15 if any(ord(ch) > 127 for ch in str(v)) else 1.0))
        if rep.totals is not None and rep.totals[j - 1] is not None:
            width = max(width, len(str(rep.totals[j - 1])) * 1.1)
        ws.column_dimensions[get_column_letter(j)].width = min(max(width + 2, 9), 48)

    ws.freeze_panes = f"A{header_row + 1}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _excel_format(kind: str) -> str:
    if kind == "money":
        return "#,##0.00"
    if kind == "percent":
        return "#,##0.0"
    if kind == "number":
        return "#,##0.###"
    return "General"


# ---------------------------------------------------------------- PDF
def _pdf_num(v, kind: str) -> str:
    """Число в «русском» виде: 1 234,56."""
    if v is None:
        return ""
    if isinstance(v, str):
        return v
    if kind == "money":
        return f"{v:,.2f}".replace(",", " ").replace(".", ",")
    if kind == "percent":
        return f"{v:,.1f}".replace(",", " ").replace(".", ",") + " %"
    if kind == "number":
        s = f"{v:,.3f}".replace(",", " ").replace(".", ",")
        if "," in s:
            s = s.rstrip("0").rstrip(",")
        return s
    return str(v)


def build_pdf(rep: Report) -> bytes:
    if not _HAS_PDF:
        raise RuntimeError("reportlab не установлен — PDF-экспорт недоступен")

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
        title=rep.title,
    )

    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph

    title_style = ParagraphStyle("t", fontName=_PDF_FONT, fontSize=14, leading=18, spaceAfter=2)
    period_style = ParagraphStyle("p", fontName=_PDF_FONT, fontSize=9, leading=12, textColor=colors.HexColor("#6B7280"))
    style = [
        Paragraph(rep.title, title_style),
    ]
    if rep.period:
        style.append(Paragraph(rep.period, period_style))

    font_size = 7.5 if len(rep.columns) > 8 else 9
    cmds = [
        ("FONTNAME", (0, 0), (-1, -1), _PDF_FONT),
        ("FONTSIZE", (0, 0), (-1, -1), font_size),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDEBF7")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9CA3AF")),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 2.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
        ("REPEATROWS", (0, 0), (-1, 0)),
    ]
    for j, c in enumerate(rep.columns):
        if c.get("align") == "right" or c["kind"] in ("number", "money", "percent"):
            cmds.append(("ALIGN", (j, 1), (j, -1), "RIGHT"))

    data = [[c["title"] for c in rep.columns]]
    for row in rep.rows:
        data.append([_pdf_num(v, rep.columns[j]["kind"]) for j, v in enumerate(row)])
    if rep.totals is not None:
        data.append([_pdf_num(v, rep.columns[j]["kind"]) for j, v in enumerate(rep.totals)])
        cmds.append(("FONTNAME", (0, len(data) - 1), (-1, len(data) - 1), _PDF_FONT))
        cmds.append(("LINEABOVE", (0, len(data) - 1), (-1, len(data) - 1), 1, colors.HexColor("#374151")))
    grid = TableStyle(cmds)

    style.append(Spacer(1, 6))
    style.append(Table(data, style=grid, repeatRows=1))
    doc.build(style)
    return buf.getvalue()
